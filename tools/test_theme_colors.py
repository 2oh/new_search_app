"""Check persisted theme fills against the production quantity-color reader."""
from pathlib import Path
import sys
import tempfile
from xml.etree import ElementTree as ET
from openpyxl import Workbook
from openpyxl.styles import Color, PatternFill
from openpyxl.writer.theme import theme_xml
from test_ui_state import load


def run(root):
    app = load('theme_app', root / 'main.py')
    with tempfile.TemporaryDirectory(prefix='theme-colors-') as temp:
        path = Path(temp) / 'colors.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.append(['数量'])
        cases = [
            (Color(theme=0), False),
            (Color(theme=0, tint=-0.05), True),
            (Color(theme=0, tint=0.5), False),
            (Color(theme=1), True),
            (Color(theme=4), True),
            (Color(theme=4, tint=0.99), True),
            (Color(theme=4, tint=1), False),
            (Color(theme=99), True),
            (Color(rgb='FFFFFFFF'), False),
            (Color(rgb='FFF2F2F2'), True),
            (Color(indexed=9), False),
            (Color(indexed=13), True),
        ]
        for row, (color, _) in enumerate(cases, 2):
            ws.cell(row, 1, 1).fill = PatternFill('solid', fgColor=color)
        ws.cell(14, 1, 1)  # no fill
        ws.merge_cells('A15:A16')
        ws['A15'].fill = PatternFill('solid', fgColor=Color(theme=0))
        ws.merge_cells('A17:A18')
        ws['A17'].fill = PatternFill('solid', fgColor=Color(theme=0, tint=-0.05))
        wb.save(path)
        expected = [value for _, value in cases] + [False, False, False, True, True]
        assert [bool(c) for c in app.get_quantity_colors(path, ws.title, 0, 0)] == expected
        # Theme slot 0 must not be assumed white; custom themes can change it.
        theme = ET.fromstring(theme_xml)
        ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
        theme.find('.//a:lt1/a:sysClr', ns).set('lastClr', 'F2F2F2')
        theme.find('.//a:accent1/a:srgbClr', ns).set('val', 'FFFFFF')
        wb.loaded_theme = ET.tostring(theme)
        wb.save(path)
        actual = [bool(c) for c in app.get_quantity_colors(path, ws.title, 0, 0)]
        assert actual[0] and not actual[4], actual
        wb.loaded_theme = '<invalid'
        wb.save(path)
        assert app.get_quantity_colors(path, ws.title, 0, 0)[0]
        wb.close()
    print('PASS: theme white, gray, black, accents, tint, custom themes, unknown themes, RGB/indexed, merged cells')


if __name__ == '__main__':
    run(Path(sys.argv[1] if len(sys.argv) > 1 else '.').resolve())
