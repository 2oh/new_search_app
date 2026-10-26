"""Regression coverage for merged quantity fills, eligibility and rendered values."""
import asyncio
import os
from pathlib import Path
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from test_ui_state import load


def texts(control):
    result = []
    if type(control).__name__ == 'Text':
        result.append(control.value)
    child = getattr(control, 'content', None)
    if child is not None:
        result.extend(texts(child))
    for child in getattr(control, 'controls', None) or []:
        result.extend(texts(child))
    return result


async def run(root):
    app = load('quantity_app', root / 'main.py')
    helper = load('quantity_helper', root / 'tools/frozen_selftest.py')
    with tempfile.TemporaryDirectory(prefix='quantity-regression-') as temp:
        path = Path(temp) / 'quantities.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = '数量確認'
        ws.append(['品番', 'PG名', '品名', '材料', '数量'])
        for i, quantity in enumerate(['各１', None, 2, 3, 0, '各１', 1.5, 0, None, 4, None], 1):
            ws.append([f'PART-{i:03}', f'PG-{i:03}', '部品', '鉄', quantity])
        yellow = PatternFill('solid', fgColor='FFFFFF00')
        ws.merge_cells('E2:E3')
        ws['E2'].fill = yellow
        ws.merge_cells('C4:C5')
        ws['E4'].fill = yellow
        ws.merge_cells('E9:E10')  # uncolored merged zero
        ws.merge_cells('E11:E12')  # uncolored merged nonzero
        wb.save(path)
        wb.close()
        colors = app.get_quantity_colors(path, ws.title, 0, 4)
        assert [bool(c) for c in colors] == [True, True, True, False, False, False, False, False, False, False, False], colors

        # Quantity column inside a rectangular merge; a non-first-row header
        # checks index offsets independently of the application's header detection.
        wb = Workbook()
        ws = wb.active
        ws['A1'] = 'title'
        ws['C2'] = '数量'
        ws.merge_cells('B3:C4')
        ws['B3'] = 2
        ws['B3'].fill = yellow
        ws.merge_cells('B5:C6')
        ws['B5'] = 3
        rectangular = Path(temp) / 'rectangular.xlsx'
        wb.save(rectangular)
        wb.close()
        assert [bool(c) for c in app.get_quantity_colors(rectangular, ws.title, 1, 2)] == [True, True, False, False]

        cwd = Path.cwd()
        os.chdir(temp)
        session = helper.Session(helper.MemoryConnection())
        session.start_updates_scheduler()
        page = session.page
        session.get_page_patch()
        token = helper._context_page.set(page)
        scope = {}
        previous_profile = sys.getprofile()
        def capture(frame, event, arg):
            if event == 'return' and frame.f_code is app.main.__code__:
                scope.update(frame.f_locals)
        try:
            sys.setprofile(capture)
            app.main(page)
            sys.setprofile(previous_profile)
            df = scope['extract_excel_dataframe'](path, '数量確認')
            assert len(df) == 11, df
            assert list(df['出力可否']) == ['不可', '不可', '不可', '可', '不可', '可', '可', '不可', '不可', '可', '可'], df
            assert list(df['数量']) == ['各１', '各１', 2, 3, 0, '各１', 1.5, 0, 0, 4, 4], df
            expected = ['各１', '各１', '2', '3', '0', '各１', '1.5', '0', '0', '4', '4']
            for width in (1200, 1920):
                page.width = width
                scope['render_table_from_df'](df)
                table = scope['table']
                column = next(i for i, c in enumerate(table.columns) if texts(c.label) == ['数量'])
                assert [texts(r.cells[column].content) for r in table.rows] == [[v] for v in expected]
                assert table.columns[column].fixed_width == 64
                assert table.data_row_height >= 72
            for value in (0, '0', '0.0', ' 0 ', '0,000'):
                assert scope['is_zero_quantity'](value)
            for value in ('各１', 2, 1.5, '', None):
                assert not scope['is_zero_quantity'](value)
            session.get_page_patch()
            print('PASS: merged/unmerged/rectangular fills, row offsets, extraction, zero exclusion, quantity text and both layouts')
        finally:
            sys.setprofile(previous_profile)
            session.close()
            helper._context_page.reset(token)
            os.chdir(cwd)


if __name__ == '__main__':
    asyncio.run(run(Path(sys.argv[1] if len(sys.argv) > 1 else '.').resolve()))
